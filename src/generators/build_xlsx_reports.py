import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from src.config import DB_PATH, EXCEL_OUTPUT_DIR

def build_xlsx_reports():
    EXCEL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    report_path = EXCEL_OUTPUT_DIR / "As_Is_Diagnostic_Report.xlsx"
    
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    wb = openpyxl.Workbook()
    
    # === PLANILHA 1: RESUMO DE ENTREVISTADOS ===
    ws_ent = wb.active
    ws_ent.title = "Matriz de Entrevistados"
    
    headers_ent = ["ID", "Nome", "Cargo", "Área", "Diretoria", "Nível", "Status Inteligência"]
    ws_ent.append(headers_ent)
    
    cursor.execute("SELECT * FROM entrevistados ORDER BY id")
    for row in cursor.fetchall():
        ws_ent.append([row['id'], row['nome'], row['cargo'], row['area'], row['diretoria'], row['nivel'], row['status_revisao']])
        
    # === PLANILHA 2: INTEGRIDADE SISTÊMICA ===
    ws_sys = wb.create_sheet(title="Uso de Sistemas e ETLs Humanos")
    headers_sys = ["ID", "Nome", "Área", "Sistema", "Satisfação", "Como Usa", "Workaround (Excel Paralelo)"]
    ws_sys.append(headers_sys)
    
    query_sys = """
        SELECT e.id, e.nome, e.area, s.sistema, s.satisfacao, s.como_usa, s.workaround 
        FROM sistemas_uso s
        JOIN entrevistados e ON s.entrevistado_id = e.id
    """
    cursor.execute(query_sys)
    for row in cursor.fetchall():
        ws_sys.append(list(row))
        
    # === PLANILHA 3: INSIGHTS DA CADEIA DE VALOR ===
    ws_val = wb.create_sheet(title="Gargalos (Cadeia de Valor)")
    headers_val = ["Etapa CAPEX", "Gravidade", "Categoria", "Sistemas Ofensores", "Relato Completo da Dor", "Citação Direta (Aspas)", "Informante"]
    ws_val.append(headers_val)
    
    query_val = """
        SELECT i.etapa_cadeia_valor, i.severidade, i.categoria, i.sistemas_envolvidos, i.descricao, i.citacao_direta, e.cargo
        FROM insights_ia i
        JOIN entrevistados e ON i.entrevistado_id = e.id
        ORDER BY i.etapa_cadeia_valor, i.severidade DESC
    """
    cursor.execute(query_val)
    for row in cursor.fetchall():
        ws_val.append(list(row))
        
    # ESTILIZANDO
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Ajusta tamanho das colunas na força bruta
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) if max_length < 50 else 50
            ws.column_dimensions[column].width = adjusted_width

    wb.save(report_path)
    conn.close()
    
    print(f"[GERADOR] [SUCESSO] Exportado Excel Consolidado em: {report_path}")

if __name__ == "__main__":
    build_xlsx_reports()
