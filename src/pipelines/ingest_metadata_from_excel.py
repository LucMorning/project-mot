import sqlite3
import openpyxl
from pathlib import Path

from src.config import DB_PATH, LISTA_ENTREVISTAS
from src.utils.matching import build_name_file_map

def ingest_master_spreadsheet():
    """Lê a planilha lista_entrevistas.xlsx e insere no banco, fazendo o link com os arquivos .docx."""
    if not LISTA_ENTREVISTAS.exists():
        print(f"Erro: Arquivo não encontrado: {LISTA_ENTREVISTAS}")
        return

    # Inicia conexão
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    # Pega o mapa de nomes => arquivo renomeado .docx
    map_files = build_name_file_map()
    
    wb = openpyxl.load_workbook(LISTA_ENTREVISTAS)
    ws = wb.active
    
    count_inserted = 0
    # Limpa tabela para testes (se quiser idempotência total pode ser merge)
    cursor.execute("DELETE FROM entrevistados")
    
    for row in range(2, ws.max_row + 1):
        nome = ws.cell(row=row, column=2).value
        # Filtra os vazios ou zeros
        if not nome or str(nome).strip() in ["", "0", "None"]:
            continue
            
        nome = str(nome).strip()
        cargo = ws.cell(row=row, column=3).value
        diretoria = ws.cell(row=row, column=4).value
        plataforma = ws.cell(row=row, column=5).value
        area = ws.cell(row=row, column=6).value
        nivel = ws.cell(row=row, column=7).value
        data_entrevista = ws.cell(row=row, column=8).value
        
        # Pode ser datetime, formata para string
        if hasattr(data_entrevista, "strftime"):
            data_entrevista = data_entrevista.strftime("%Y-%m-%d %H:%M")
        elif data_entrevista:
            data_entrevista = str(data_entrevista)
            
        tipo_entrevista = ws.cell(row=row, column=9).value
        
        # Pega a lista de arquivos disponiveis no novo diretorio
        from src.config import TRANSCRICOES_DIR
        import os
        from src.utils.matching import get_best_match
        
        available = os.listdir(TRANSCRICOES_DIR) if TRANSCRICOES_DIR.exists() else []
        arquivo_docx = get_best_match(nome, available)
        
        cursor.execute("""
            INSERT INTO entrevistados 
            (nome, cargo, diretoria, plataforma, area, nivel, data_entrevista, tipo_entrevista, arquivo_transcricao)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (nome, cargo, diretoria, plataforma, area, nivel, data_entrevista, tipo_entrevista, arquivo_docx))
        
        count_inserted += 1

    # Lógica QA Integridade: Inserir DOCXs físicos que sobraram e não estão no Excel
    used_files = [x[0] for x in cursor.execute("SELECT arquivo_transcricao FROM entrevistados WHERE arquivo_transcricao IS NOT NULL").fetchall()]
    for arquivo_solto in available:
        if arquivo_solto not in used_files:
            nome_assumed = arquivo_solto.replace("entrevista_as_is_", "").replace("aprofundamento_", "").replace(".docx", "").replace("_", " ").title()
            cursor.execute("""
            INSERT INTO entrevistados 
            (nome, cargo, arquivo_transcricao, tipo_entrevista, status_revisao)
            VALUES (?, 'Não Identificado (Arquivo Extra)', ?, 'Extra QA', 'pendente')
            """, (nome_assumed, arquivo_solto))
            count_inserted += 1

    conn.commit()
    conn.close()
    print(f"[SUCESSO] Ingestão completa! {count_inserted} entrevistados carregados no banco ({DB_PATH.name}).")

if __name__ == "__main__":
    ingest_master_spreadsheet()
