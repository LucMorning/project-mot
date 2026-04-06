"""
Ingest Metadata - ETL da planilha mestre para stg_entrevistados.

Responsabilidade unica: ler lista_entrevistas.xlsx e popular stg_entrevistados.

NOTA: Arquivos de transcrição são gerenciados em stg_arquivos (extract_transcripts).
"""
import re
import os
import sqlite3
import openpyxl
import unidecode as _unidecode

from src.config import (
    DB_PATH, LISTA_ENTREVISTAS, TRANSCRICOES_DIR,
    TRANSCRIPT_PREFIXES, TRANSCRIPT_EXTENSION,
    normalize_unidade,
)
from src.utils.name_matcher import build_name_file_map, get_best_match
from src.workflows.ingest.deduplicator import (
    cleanup_duplicate_interviewees,
    find_existing_interviewee,
)


def ingest_master_spreadsheet():
    """
    Lê a planilha lista_entrevistas.xlsx e insere em stg_entrevistados.

    Processo:
    1. Limpa duplicatas existentes
    2. Insere/atualiza dados do Excel
    """
    if not LISTA_ENTREVISTAS.exists():
        print(f"Erro: Arquivo não encontrado: {LISTA_ENTREVISTAS}")
        return

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    # 1. LIMPA DUPLICATAS ANTES DE INGERIR
    print("\n[CLEANUP] Removendo duplicatas existentes...")
    stats_before = cleanup_duplicate_interviewees(conn)
    print(f"[CLEANUP] {stats_before['removed']} duplicatas removidas.\n")

    # Pega o mapa de nomes => arquivo renomeado .docx
    map_files = build_name_file_map()

    wb = openpyxl.load_workbook(LISTA_ENTREVISTAS)
    ws = wb.active

    count_inserted = 0
    count_updated = 0

    for row in range(2, ws.max_row + 1):
        nome = ws.cell(row=row, column=2).value
        # Filtra os vazios ou zeros
        if not nome or str(nome).strip() in ["", "0", "None"]:
            continue

        # Nome: UPPERCASE + sem acentos
        nome = _unidecode.unidecode(str(nome).strip()).upper()

        # Remove sufixos como " - Parte 2", " - A", "PARTE 2"
        nome = re.sub(r'\s*-\s*Parte\s+\d+', '', nome, flags=re.IGNORECASE)
        nome = re.sub(r'\s+Parte\s+\d+$', '', nome, flags=re.IGNORECASE)
        nome = re.sub(r'\s*-\s*[A-Z]$', '', nome)

        cargo = ws.cell(row=row, column=3).value
        diretoria = ws.cell(row=row, column=4).value
        unidade_negocio = normalize_unidade(ws.cell(row=row, column=5).value)
        area = ws.cell(row=row, column=6).value
        nivel = ws.cell(row=row, column=7).value
        dt_entrevista = ws.cell(row=row, column=8).value

        # Diretoria: NULL se for "0", "#N/A", vazio
        if diretoria and str(diretoria).strip() in ["0", "#N/A", "", "None", "NA"]:
            diretoria = None

        # Pode ser datetime, formata para string
        if hasattr(dt_entrevista, "strftime"):
            dt_entrevista = dt_entrevista.strftime("%Y-%m-%d %H:%M")
        elif dt_entrevista:
            dt_entrevista = str(dt_entrevista)

        tipo_entrevista = ws.cell(row=row, column=9).value

        # Verifica se já existe (UPDATE ou INSERT)
        existing_id = find_existing_interviewee(conn, nome)

        if existing_id:
            # UPDATE do existente - usa o nome mais longo (mais completo)
            cursor.execute("SELECT nome FROM stg_entrevistados WHERE id = ?", (existing_id,))
            current_nome = cursor.fetchone()[0]
            nome_to_use = nome if len(nome) > len(current_nome) else current_nome

            cursor.execute("""
                UPDATE stg_entrevistados
                SET nome = ?, cargo = ?, diretoria = ?, unidade_negocio = ?, area = ?, nivel = ?,
                    dt_entrevista = ?, tipo_entrevista = ?
                WHERE id = ?
            """, (nome_to_use, cargo, diretoria, unidade_negocio, area, nivel, dt_entrevista, tipo_entrevista, existing_id))
            count_updated += 1
        else:
            # INSERT novo
            cursor.execute("""
                INSERT INTO stg_entrevistados
                (nome, cargo, diretoria, unidade_negocio, area, nivel, dt_entrevista, tipo_entrevista)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (nome, cargo, diretoria, unidade_negocio, area, nivel, dt_entrevista, tipo_entrevista))
            count_inserted += 1

    # 2. LIMPA DUPLICATAS DEPOIS DE INGERIR
    print("\n[CLEANUP] Removendo duplicatas criadas pelo Excel...")
    stats_after = cleanup_duplicate_interviewees(conn)
    print(f"[CLEANUP] {stats_after['removed']} duplicatas removidas.\n")

    conn.commit()
    conn.close()
    total_removed = stats_before['removed'] + stats_after['removed']
    print(f"[SUCESSO] Ingestão completa! {total_removed} duplicatas removidas, {count_inserted} novos, {count_updated} atualizados.")


def main():
    """Entry point para Poetry scripts."""
    ingest_master_spreadsheet()


if __name__ == "__main__":
    main()
