from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from openpyxl import load_workbook
import os

def create_strategic_pptx():
    print("Gerando apresentação estratégica de Engenharia de Dados (PPTX)...")
    prs = Presentation()

    # Estilos de cores (Simulando identidade visual sóbria e técnica)
    # Azul Marinho: 31, 78, 120
    # Vermelho Alerta: 192, 0, 0

    def add_title_slide(title_text, subtitle_text):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title_text
        slide.placeholders[1].text = subtitle_text

    def add_bullet_slide(title_text, items, level=0):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title_text
        tf = slide.placeholders[1].text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = item
            p.level = level

    # 1. Título
    add_title_slide("Diagnóstico de CAPEX: Processos e Dados", 
                    "Visão Estratégica de Engenharia de Dados\nProjeto MOTIVA - Março 2026")

    # 2. Nossa Lente
    add_bullet_slide("Nossa Lente: Engenharia de Dados Sênior", [
        "Foco na Linhagem do Dado (Data Lineage): Rastreabilidade do campo ao report executivo.",
        "Identificação de 'ETLs Humanos': Processos manuais de extração, tratamento e carga.",
        "Diagnóstico de Interoperabilidade: Por que os 25+ sistemas não se comunicam?",
        "Integridade Sistêmica: Redução de planilhas 'sombra' e silos de informação."
    ])

    # 3. Cadeia de Valor de TI
    add_bullet_slide("Cadeia de Valor do CAPEX (7 Etapas)", [
        "1. Novos Negócios & Demandas",
        "2. Orçamento (Budgeting)",
        "3. Estruturação (Planejamento Físico-Financeiro)",
        "4. Contratação & Execução (Compras/Aditivos)",
        "5. Medição (RDO e Boletins)",
        "6. Tendência (Forecast e Riscos)",
        "7. Fiscal & Pagamento (NF e Conciliação)"
    ])

    # 4. O Ecossistema Fragmentado (As-Is)
    add_bullet_slide("O Labirinto Sistêmico (25+ Aplicações)", [
        "Core: SAP (PS/FI/BW/BPC), Archer, SIC.",
        "Engenharia: P6, Prisma, Compor 90, Kartado, Fulcrum.",
        "Suporte: Coupa, Netlex, Flexchain, Docusign, V360, Atlas.",
        "Sombra: Uso massivo de Excel, SharePoint e Forms como integradores manuais."
    ])

    # 5. Diagnóstico de 'Doenças' Sistêmicas
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Diagnóstico: As 'Doenças' de Dados"
    tf = slide.placeholders[1].text_frame
    
    doencas = [
        ("Quebra de Linhagem", "O dado de medição física no campo (Kartado/RDO) não flui automaticamente para a medição financeira no SAP."),
        ("Silos de Informação", "Áreas como 'Riscos' (Archer) e 'Engenharia' (P6) operam com bases de dados dessincronizadas."),
        ("Latência Operacional", "Aprovoções dependem de 'Workflows' manuais via e-mail e Teams, fora das trilhas de auditoria sistêmica."),
        ("Risco de Governança", "A falta de centralização de memórias de cálculo em projetos de R$ 8 bi aumenta o risco de pleitos indevidos.")
    ]
    
    for titulo, desc in doencas:
        p = tf.add_paragraph()
        p.text = titulo
        p.font.bold = True
        p.font.size = Pt(18)
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.level = 1
        p_desc.font.size = Pt(14)

    # 6. Evidências Cruas (Extraídas das 68 Entrevistas)
    if os.path.exists("diagnostico_estrategico_ti_processos.xlsx"):
        wb = load_workbook("diagnostico_estrategico_ti_processos.xlsx")
        ws = wb.active
        
        for row in ws.iter_rows(min_row=2, max_row=8, values_only=True):
            etapa = row[0]
            dores = row[2]
            if dores and len(dores) > 20:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = f"Detalhamento: {etapa}"
                body = slide.placeholders[1]
                tf = body.text_frame
                tf.word_wrap = True
                
                p = tf.add_paragraph()
                p.text = "Evidências e Falhas Encontradas:"
                p.font.bold = True
                
                # Split dores by bullets and add as paragraphs
                for d in dores.split("•"):
                    if d.strip():
                        p_dor = tf.add_paragraph()
                        p_dor.text = d.strip()
                        p_dor.level = 1
                        p_dor.font.size = Pt(12)

    # 7. Conclusão do As-Is
    add_bullet_slide("Conclusão: Impacto no Negócio", [
        "A fragmentação atual impede uma 'Single Source of Truth'.",
        "O retrabalho em 'ETLs manuais' consome horas técnicas valiosas de engenharia.",
        "Existe um gap crítico de visibilidade de Tendência (Forecast) em tempo real.",
        "O diagnóstico micro prova que o problema não é a falta de sistemas, mas a falta de CONEXÃO entre eles."
    ])

    # Salvar
    output_file = "DIAGNOSTICO_ESTRATEGICO_CAPEX_MOTIVA.pptx"
    prs.save(output_file)
    print(f"Relatório estratégico gerado: {output_file}")

if __name__ == "__main__":
    create_strategic_pptx()
