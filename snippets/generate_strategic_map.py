import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import zipfile
import xml.etree.ElementTree as ET
import os
import re

# CONFIGURAÇÕES
TRANSCRICOES_DIR = "05. Transcrições Integrais"
LISTA_REF = "lista de entrevistas.xlsx"
OUTPUT_XLSX = "diagnostico_estrategico_ti_processos.xlsx"

# 7 ETAPAS DA CADEIA DE VALOR (CONFORME RELAÓRIO TI RÔMULO)
CADEIA_VALOR = {
    "1. Novos Negócios & Demandas": ["abertura", "demanda", "aprovacao", "viabilidade", "novo negócio", "proposta"],
    "2. Orçamento": ["orçamento", "budget", "custos", "capex", "estimativa", "valor"],
    "3. Estruturação (Plan/Custo)": ["planejamento", "cronograma", "eap", "wbs", "físico", "financeiro", "compor 90", "prisma", "project"],
    "4. Contratação & Execução": ["contrato", "aditivo", "compras", "coupa", "netlex", "flexchain", "docusign", "suprimentos"],
    "5. Medição": ["medição", "boletim", "rdo", "relatório diário", "conecta", "fulcrum", "kartado", "validação"],
    "6. Tendência": ["tendência", "projeção", "forecast", "risco", "archer", "bw", "bpc", "power bi"],
    "7. Fiscal (NF) & Pagamento": ["nota fiscal", "nf", "pagamento", "v360", "atlas", "conciliação", "fi", "ap"]
}

SISTEMAS_LIST = [
    "SAP PS", "SAP FI", "SAP BW", "SAP BPC", "Prisma", "Project", "Excel", "Compor 90", 
    "Coupa", "Netlex", "Flexchain", "Docusign", "Teams", "Kartado", "SharePoint", 
    "Power Apps", "PM", "Conecta", "Fulcrum", "Archer", "V360", "Atlas", "Forms"
]

DORES_KEYWORDS = ["manual", "duplicidade", "erro", "demora", "não integra", "planilha", "paralelo", "sombra", "falta", "difícil", "retrabalho"]

def extract_text_from_docx(file_path):
    if not os.path.exists(file_path): return ""
    text = []
    try:
        with zipfile.ZipFile(file_path, 'r') as zip_ref:
            with zip_ref.open('word/document.xml') as f:
                tree = ET.parse(f)
                root = tree.getroot()
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for p in root.findall('.//w:p', ns):
                    para = "".join([t.text for t in p.findall('.//w:t', ns) if t.text])
                    if para: text.append(para)
    except: return ""
    return "\n".join(text)

def run_strategic_diagnostic():
    print("Iniciando Diagnóstico Estratégico baseado na Cadeia de Valor de TI...")
    
    # Mapeamento de Pessoas (congelado do Excel de referência)
    mapping_people = {}
    try:
        wb_ref = openpyxl.load_workbook(LISTA_REF)
        ws_ref = wb_ref.active
        for r in range(2, ws_ref.max_row + 1):
            name = ws_ref.cell(row=r, column=2).value
            area = ws_ref.cell(row=r, column=6).value
            cargo = ws_ref.cell(row=r, column=3).value
            if name: mapping_people[str(name).lower().strip()] = f"{cargo} ({area})"
    except: pass

    # Resultados agrupados por Etapa da Cadeia de Valor
    etapas_data = {etapa: {"insights": [], "sistemas": set(), "dores": set(), "fontes": set()} for etapa in CADEIA_VALOR}
    
    files = [f for f in os.listdir(TRANSCRICOES_DIR) if f.endswith('.docx')]
    
    for filename in files:
        path = os.path.join(TRANSCRICOES_DIR, filename)
        text = extract_text_from_docx(path)
        if not text: continue
        
        # Identificar Fonte
        autor = "Desconhecido"
        for p_name, p_info in mapping_people.items():
            name_parts = p_name.split()
            if len(name_parts) >= 2:
                if name_parts[0] in filename.lower() and name_parts[-1] in filename.lower():
                    autor = p_info
                    break
        
        paragraphs = text.split("\n")
        for para in paragraphs:
            para_lower = para.lower()
            if len(para) < 40: continue

            for etapa, kws in CADEIA_VALOR.items():
                if any(kw in para_lower for kw in kws):
                    # Identificou o contexto da etapa
                    etapas_data[etapa]["fontes"].add(autor)
                    
                    # Checar Sistemas
                    for s in SISTEMAS_LIST:
                        if s.lower() in para_lower:
                            etapas_data[etapa]["sistemas"].add(s)
                    
                    # Checar Dores
                    if any(dkw in para_lower for dkw in DORES_KEYWORDS):
                        etapas_data[etapa]["dores"].add(para.strip())
                    
                    # Guardar insight geral
                    etapas_data[etapa]["insights"].append(para.strip())

    # CRIAR EXCEL FINAL
    wb = Workbook()
    ws = wb.active
    ws.title = "Cadeia de Valor As-Is"
    
    headers = ["Etapa da Cadeia de Valor", "Cenário Atual (Detalhado)", "Dores e Falhas de Integração", "Sistemas Envolvidos", "Áreas Contribuintes"]
    
    # Estilos
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    top_align = Alignment(vertical="top", wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for col, head in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = 2
    for etapa, data in etapas_data.items():
        if not data["insights"] and not data["dores"]: continue
        
        # Consolida tudo sem resumos ou cortes
        sintese = "\n\n".join(list(dict.fromkeys(data["insights"]))) # Remove duplicatas exatas
        dores = "\n".join([f"• {d}" for d in list(data["dores"])])
        sistemas = ", ".join(sorted(list(data["sistemas"])))
        fontes = ", ".join(sorted(list(data["fontes"])))

        ws.cell(row=row_idx, column=1, value=etapa).alignment = top_align
        ws.cell(row=row_idx, column=2, value=sintese).alignment = top_align
        ws.cell(row=row_idx, column=3, value=dores).alignment = top_align
        ws.cell(row=row_idx, column=4, value=sistemas).alignment = top_align
        ws.cell(row=row_idx, column=5, value=fontes).alignment = top_align
        
        for c in range(1, 6): ws.cell(row=row_idx, column=c).border = border
        row_idx += 1

    # Ajuste de Layout
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 70
    ws.column_dimensions['C'].width = 70
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 40
    
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUTPUT_XLSX)
    print(f"Relatório gerado: {OUTPUT_XLSX}")

if __name__ == "__main__":
    run_strategic_diagnostic()
